"""Per-day attribution: join daily usage exports (one CSV per day) with portal login events."""
import re, io, collections
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from merge import parse_usage, LOGIN_RE, MAC_RE, dev_label

DATE_RE = re.compile(r'(\d{4}-\d{2}-\d{2})')


def date_from_name(name):
    m = DATE_RE.search(name or '')
    return date.fromisoformat(m.group(1)) if m else None


def dates_from_name(name):
    """All YYYY-MM-DD dates in a file name, in order (a 7-day export carries a start and an end date)."""
    out = []
    for d in DATE_RE.findall(name or ''):
        try:
            out.append(date.fromisoformat(d))
        except ValueError:
            pass
    return out


def parse_events_by_day(fileobj):
    """Return {mac: [(datetime, user), ...]} sorted ascending (successful logins only)."""
    wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h else '' for h in next(rows)]
    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None
    i_dt, i_mac, i_det = col('date/time', 'date', 'time'), col('client mac', 'mac'), col('details', 'event', 'message')
    if i_dt is None or i_mac is None or i_det is None:
        raise ValueError('Event log must have "Date/Time", "Client MAC" and "Details" columns')
    ev = collections.defaultdict(list)
    for r in rows:
        dt, mac, det = r[i_dt], r[i_mac], r[i_det]
        if not (dt and mac and det):
            continue
        m = LOGIN_RE.search(str(det))
        if not m:
            continue
        if isinstance(dt, str):
            try:
                dt = datetime.fromisoformat(dt)
            except ValueError:
                continue
        ev[str(mac).strip().upper().replace('-', ':')].append((dt, m.group(1)))
    for k in ev:
        ev[k].sort()
    return ev


def attribute(ev_list, day):
    """Pick user for a MAC on a given day.
    Returns (user, basis, users_that_day_str). basis: 'login' | 'carried' | 'inferred' | ''"""
    same = [u for dt, u in ev_list if dt.date() == day]
    if same:
        c = collections.Counter(same)
        return c.most_common(1)[0][0], 'login', ', '.join(f'{u} ({n})' for u, n in c.most_common())
    before = [(dt, u) for dt, u in ev_list if dt.date() < day]
    if before:
        return before[-1][1], 'carried', ''
    after = [(dt, u) for dt, u in ev_list if dt.date() > day]
    if after:  # event log starts after this day (export truncated) - use the next known login
        return after[0][1], 'inferred', ''
    return '', '', ''


def merge_daily(event_fileobj, usage_files):
    """usage_files: list of (filename, fileobj). Returns (detail_rows, pivot, days, summary)."""
    ev = parse_events_by_day(event_fileobj)
    byday = []
    for name, fobj in usage_files:
        d = date_from_name(name)
        if not d:
            raise ValueError(f'Cannot find a date (YYYY-MM-DD) in file name: {name}')
        byday.append((d, parse_usage(fobj)))
    return compute(ev, byday)


def compute(ev, usage_by_day):
    """ev: {mac: [(dt,user)...]} sorted; usage_by_day: [(date, {mac: usage dict})]."""
    days = []
    detail = []
    for d, usage in usage_by_day:
        days.append(d)
        for mac, u in usage.items():
            user, basis, seen = attribute(ev.get(mac, []), d) if u['is_mac'] else ('', '', '')
            detail.append({'day': d, 'mac': mac, 'name': u['name'], 'user': user, 'basis': basis, 'seen': seen,
                           'up': round(u['up'], 2), 'down': round(u['down'], 2), 'tot': round(u['tot'], 2),
                           'is_mac': u['is_mac']})
    days = sorted(set(days))
    detail.sort(key=lambda r: (r['day'], -r['tot']))
    per = collections.defaultdict(lambda: {m: {d: 0.0 for d in days} for m in ('tot', 'up', 'down')})
    macs = collections.defaultdict(set)
    names = {}
    for r in detail:
        if r['user']:
            for m in ('tot', 'up', 'down'):
                per[r['user']][m][r['day']] += r[m]
            macs[r['user']].add(r['mac'])
            if r['name']:
                names[r['mac']] = r['name']
    pivot_rows = []
    for u, p in per.items():
        total, up, down = sum(p['tot'].values()), sum(p['up'].values()), sum(p['down'].values())
        pivot_rows.append({'user': u, 'per': p['tot'], 'per_up': p['up'], 'per_down': p['down'],
                           'total': round(total, 2), 'up': round(up, 2), 'down': round(down, 2), 'gb': round(total / 1024, 2),
                           'macs': sorted(macs[u]), 'devs': [dev_label(m, names.get(m, '')) for m in sorted(macs[u])],
                           'days_active': sum(1 for v in p['tot'].values() if v > 0)})
    pivot_rows.sort(key=lambda r: -r['total'])
    unmatched_tot = sum(r['tot'] for r in detail if not r['user'])
    summary = {'days': len(days), 'first': days[0] if days else None, 'last': days[-1] if days else None,
               'users': len(pivot_rows), 'device_days': len(detail),
               'matched': sum(1 for r in detail if r['user']),
               'carried': sum(1 for r in detail if r['basis'] == 'carried'),
               'inferred': sum(1 for r in detail if r['basis'] == 'inferred'),
               'unmatched_gb': round(unmatched_tot / 1024, 2),
               'up_gb': round(sum(r['up'] for r in detail) / 1024, 2),
               'down_gb': round(sum(r['down'] for r in detail) / 1024, 2),
               'total_gb': round(sum(r['tot'] for r in detail) / 1024, 2)}
    return detail, pivot_rows, days, summary


def device_pivot(detail, days):
    """One row per device (MAC) with usage per day: [{mac, name, users, per/per_up/per_down{day: MB}, total, up, down, gb,
    days_active}], biggest first. users = every account the device was attributed to over the range, in order of first appearance."""
    devs = collections.OrderedDict()
    for r in detail:                      # detail is sorted by day, so first appearance = chronological
        d = devs.setdefault(r['mac'], {'mac': r['mac'], 'name': '', 'users': [], 'per': {x: 0.0 for x in days},
                                       'per_up': {x: 0.0 for x in days}, 'per_down': {x: 0.0 for x in days}})
        if r['name']:
            d['name'] = r['name']
        if r['user'] and r['user'] not in d['users']:
            d['users'].append(r['user'])
        d['per'][r['day']] += r['tot']
        d['per_up'][r['day']] += r['up']
        d['per_down'][r['day']] += r['down']
    rows = []
    for d in devs.values():
        total = sum(d['per'].values())
        d.update(total=round(total, 2), up=round(sum(d['per_up'].values()), 2), down=round(sum(d['per_down'].values()), 2),
                 gb=round(total / 1024, 2), days_active=sum(1 for v in d['per'].values() if v > 0))
        rows.append(d)
    rows.sort(key=lambda r: -r['total'])
    return rows


def unmatched_devices(detail):
    """Devices (whole range) that never got a user: [{mac, name, up, down, tot, gb, note}], biggest first."""
    out = collections.OrderedDict()
    for r in detail:
        if r['user']:
            continue
        d = out.setdefault(r['mac'], {'mac': r['mac'], 'name': '', 'up': 0.0, 'down': 0.0, 'tot': 0.0,
                                      'note': 'No login found' if r['is_mac'] else 'Not a MAC (IP / non-portal traffic)'})
        d['name'] = r['name'] or d['name']
        d['up'] += r['up']; d['down'] += r['down']; d['tot'] += r['tot']
    rows = list(out.values())
    for d in rows:
        d.update(up=round(d['up'], 2), down=round(d['down'], 2), tot=round(d['tot'], 2), gb=round(d['tot'] / 1024, 2))
    rows.sort(key=lambda r: -r['tot'])
    return rows


def compute_period(ev, usage, start, end):
    """One usage file covering start..end (e.g. 7 days): one row per device, user = most logins inside the period;
    else last earlier login (carried); else next later login (inferred). Same row shape as merge.merge()."""
    rows = []
    logged = set()
    for mac, u in usage.items():
        evl = ev.get(mac, []) if u['is_mac'] else []
        inside = [(dt, usr) for dt, usr in evl if start <= dt.date() <= end]
        user = users = lt = lu = basis = ''
        n = 0
        if inside:
            c = collections.Counter(usr for _, usr in inside)
            user = c.most_common(1)[0][0]
            users = ', '.join(f'{x} ({k})' for x, k in c.most_common())
            lt, lu, n, basis = inside[-1][0].strftime('%Y-%m-%d %H:%M'), inside[-1][1], len(inside), 'login'
        else:
            before = [x for x in evl if x[0].date() < start]
            after = [x for x in evl if x[0].date() > end]
            if before:
                user, users, basis = before[-1][1], '(carried from earlier login)', 'carried'
            elif after:
                user, users, basis = after[0][1], '(inferred from next login)', 'inferred'
        if evl:
            logged.add(mac)
        rows.append({'user': user, 'mac': mac, 'name': u['name'], 'up': round(u['up'], 2), 'down': round(u['down'], 2),
                     'tot': round(u['tot'], 2), 'gb': round(u['tot'] / 1024, 2), 'logins': n, 'users': users,
                     'last_user': lu, 'last_time': lt, 'pct': round(u.get('pct', 0.0), 2), 'is_mac': u['is_mac'], 'basis': basis})
    rows.sort(key=lambda r: -r['tot'])
    per_user = collections.OrderedDict()
    for r in rows:
        if r['user']:
            p = per_user.setdefault(r['user'], {'macs': [], 'devs': [], 'tot': 0.0})
            p['macs'].append(r['mac'])
            p['devs'].append(dev_label(r['mac'], r['name']))
            p['tot'] += r['tot']
    per_user = collections.OrderedDict(sorted(per_user.items(), key=lambda x: -x[1]['tot']))
    unmatched = [r for r in rows if not r['user']]
    summary = {'non_mac': sum(1 for r in rows if not r['is_mac']), 'devices': len(rows),
               'matched': len(rows) - len(unmatched), 'unmatched': len(unmatched), 'users': len(per_user),
               'login_macs': len(logged), 'total_gb': round(sum(r['tot'] for r in rows) / 1024, 2)}
    return rows, per_user, unmatched, summary


BASIS_TEXT = {'login': 'Login in period', 'carried': 'Carried from earlier login', 'inferred': 'Inferred from next login (log starts later)'}


def _fill_period_device_sheet(sh, rows):
    sh.append(['MAC address', 'Hostname', 'User ID', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Total (GB)', '% Usage', 'Basis'])
    for r in rows:
        sh.append([r['mac'], r['name'], r['user'], r['up'], r['down'], r['tot'], r['gb'], r['pct'],
                   BASIS_TEXT.get(r.get('basis'), 'No login found' if r['is_mac'] else 'Not a MAC')])
    sh.append(['TOTAL', '', '', round(sum(r['up'] for r in rows), 2), round(sum(r['down'] for r in rows), 2),
               round(sum(r['tot'] for r in rows), 2), round(sum(r['tot'] for r in rows) / 1024, 2), '', ''])
    for c in sh[sh.max_row]:
        c.font = Font(bold=True)
    _style(sh)


def build_period_device_xlsx(rows):
    """Per-device export of a period / single-file merge on its own (one sheet)."""
    out = openpyxl.Workbook()
    sh = out.active
    sh.title = 'Usage per device'
    _fill_period_device_sheet(sh, rows)
    buf = io.BytesIO()
    out.save(buf)
    buf.seek(0)
    return buf


def build_period_xlsx(rows, per_user, unmatched):
    """Period (e.g. 7-day) export: the usual 3 sheets from merge.build_xlsx plus 'Usage per device' first."""
    from merge import build_xlsx
    wb = openpyxl.load_workbook(build_xlsx(rows, per_user, unmatched))
    _fill_period_device_sheet(wb.create_sheet('Usage per device', 0), rows)
    _style(sh)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _style(sh):
    for c in sh[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1F4E78')
    sh.freeze_panes = 'B2'
    sh.auto_filter.ref = sh.dimensions
    for col in sh.columns:
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        sh.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, w + 2), 50)
    for row in sh.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00'


METRICS = (('tot', 'per', '', 'Total'), ('up', 'per_up', ' - Upload', 'Upload'), ('down', 'per_down', ' - Download', 'Download'))


def _fill_device_sheet(sd, devs, days, key='per'):
    """Per-device-per-day sheet for one metric (key: per | per_up | per_down)."""
    tot_key = {'per': 'total', 'per_up': 'up', 'per_down': 'down'}[key]
    sd.append(['MAC address', 'Hostname', 'User ID'] + [d.strftime('%d %b') for d in days] + [f'{dict(per="Total", per_up="Upload", per_down="Download")[key]} (MB)', 'Total (GB)' if key == 'per' else '', 'Days active'])
    for r in devs:
        sd.append([r['mac'], r['name'], ', '.join(r['users'])] + [round(r[key][d], 2) for d in days] + [r[tot_key], r['gb'] if key == 'per' else '', r['days_active']])
    sd.append(['TOTAL', '', ''] + [round(sum(r[key][d] for r in devs), 2) for d in days] + [round(sum(r[tot_key] for r in devs), 2), '', ''])
    for c in sd[sd.max_row]:
        c.font = Font(bold=True)


def _fill_user_day_sheet(s, pivot_rows, days, key='per'):
    tot_key = {'per': 'total', 'per_up': 'up', 'per_down': 'down'}[key]
    label = {'per': 'Total', 'per_up': 'Upload', 'per_down': 'Download'}[key]
    s.append(['User ID'] + [d.strftime('%d %b') for d in days] + [f'{label} (MB)', 'Total (GB)' if key == 'per' else '', 'Days active', 'No. of devices', 'MACs (hostname)'])
    for r in pivot_rows:
        s.append([r['user']] + [round(r[key][d], 2) for d in days] + [r[tot_key], r['gb'] if key == 'per' else '', r['days_active'], len(r['macs']), ', '.join(r['devs'])])
    s.append(['TOTAL'] + [round(sum(r[key][d] for r in pivot_rows), 2) for d in days] + [round(sum(r[tot_key] for r in pivot_rows), 2), '', '', '', ''])
    for c in s[s.max_row]:
        c.font = Font(bold=True)


def _save(wb):
    for sh in wb:
        _style(sh)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_device_xlsx(detail, days, view='day'):
    """Per-device export on its own. view 'total' = one row per device for the range (upload / download / total);
    'day' = per-day sheets for Total, Upload and Download."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    devs = device_pivot(detail, days)
    if view == 'total':
        _fill_total_device_sheet(wb.create_sheet('Usage per device'), devs)
    else:
        for _, key, suffix, _ in METRICS:
            _fill_device_sheet(wb.create_sheet('Usage per device per day' if not suffix else 'Device per day' + suffix), devs, days, key)
    return _save(wb)


def _fill_total_device_sheet(sd, devs):
    sd.append(['MAC address', 'Hostname', 'User ID', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Total (GB)', 'Days active'])
    for r in devs:
        sd.append([r['mac'], r['name'], ', '.join(r['users']), r['up'], r['down'], r['total'], r['gb'], r['days_active']])
    sd.append(['TOTAL', '', '', round(sum(r['up'] for r in devs), 2), round(sum(r['down'] for r in devs), 2),
               round(sum(r['total'] for r in devs), 2), round(sum(r['total'] for r in devs) / 1024, 2), ''])
    for c in sd[sd.max_row]:
        c.font = Font(bold=True)


def build_total_xlsx(detail, pivot_rows, days):
    """Total-period workbook: per user, per device (both with upload / download / total) and unmatched."""
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = 'Usage per user'
    s.append(['User ID', 'No. of devices', 'MACs (hostname)', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Total (GB)'])
    for r in pivot_rows:
        s.append([r['user'], len(r['macs']), ', '.join(r['devs']), r['up'], r['down'], r['total'], r['gb']])
    s.append(['TOTAL', '', '', round(sum(r['up'] for r in pivot_rows), 2), round(sum(r['down'] for r in pivot_rows), 2),
              round(sum(r['total'] for r in pivot_rows), 2), round(sum(r['total'] for r in pivot_rows) / 1024, 2)])
    for c in s[s.max_row]:
        c.font = Font(bold=True)
    _fill_total_device_sheet(wb.create_sheet('Usage per device'), device_pivot(detail, days))
    u = wb.create_sheet('Unmatched')
    u.append(['MAC / ID', 'Hostname', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Note'])
    for r in unmatched_devices(detail):
        u.append([r['mac'], r['name'], r['up'], r['down'], r['tot'], r['note']])
    return _save(wb)


def build_daily_xlsx(detail, pivot_rows, days):
    """Day-wise workbook: per user and per device for Total, Upload and Download, device-day detail, unmatched."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    devs = device_pivot(detail, days)
    for _, key, suffix, _ in METRICS:
        _fill_user_day_sheet(wb.create_sheet('Usage per user per day' if not suffix else 'User per day' + suffix), pivot_rows, days, key)
    for _, key, suffix, _ in METRICS:
        _fill_device_sheet(wb.create_sheet('Usage per device per day' if not suffix else 'Device per day' + suffix), devs, days, key)
    s2 = wb.create_sheet('Device-day detail')
    s2.append(['Date', 'User ID', 'MAC address', 'Hostname', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Basis', 'Users logged in that day (logins)'])
    for r in detail:
        s2.append([r['day'].isoformat(), r['user'], r['mac'], r['name'], r['up'], r['down'], r['tot'],
                   {'login': 'Login same day', 'carried': 'Carried from earlier login', 'inferred': 'Inferred from next login (log starts later)', '': ('No login found' if r['is_mac'] else 'Not a MAC')}[r['basis']],
                   r['seen']])
    s3 = wb.create_sheet('Unmatched')
    s3.append(['Date', 'MAC / ID', 'Hostname', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Note'])
    for r in detail:
        if not r['user']:
            s3.append([r['day'].isoformat(), r['mac'], r['name'], r['up'], r['down'], r['tot'], 'No login found' if r['is_mac'] else 'Not a MAC (IP / non-portal traffic)'])
    return _save(wb)

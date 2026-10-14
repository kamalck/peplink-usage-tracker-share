"""Core logic: join Starlink portal login events (xlsx) with per-device usage (csv)."""
import csv, io, re, collections
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MAC_RE = re.compile(r'^([0-9A-F]{2}:){5}[0-9A-F]{2}$')
LOGIN_RE = re.compile(r'guest account ID "([^"]+)" \[login\]')


def dev_label(mac, name):
    """'MAC (hostname)' - the hostname is the Name column of the usage export."""
    return f'{mac} ({name})' if name else mac


def parse_event_log(fileobj):
    """Return (logins{mac: Counter(user)}, last{mac: (dt, user)})."""
    wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h else '' for h in next(rows)]
    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None
    i_dt, i_mac, i_det = col('date/time', 'date', 'time'), col('client mac', 'mac'), col('details', 'event', 'message')
    if i_mac is None or i_det is None:
        raise ValueError('Event log must have "Client MAC" and "Details" columns')
    logins = collections.defaultdict(collections.Counter)
    last = {}
    for r in rows:
        mac, det = r[i_mac], r[i_det]
        if not mac or not det:
            continue
        m = LOGIN_RE.search(str(det))
        if not m:
            continue
        mac = str(mac).strip().upper().replace('-', ':')
        user = m.group(1)
        logins[mac][user] += 1
        dt = r[i_dt] if i_dt is not None else None
        if isinstance(dt, str):
            try:
                dt = datetime.fromisoformat(dt)
            except ValueError:
                dt = None
        if dt and (mac not in last or dt > last[mac][0]):
            last[mac] = (dt, user)
    return logins, last


def parse_usage(fileobj):
    """Return OrderedDict{mac: {name, up, down, tot}} (duplicate MACs summed)."""
    text = fileobj.read()
    if isinstance(text, bytes):
        text = text.decode('utf-8-sig', errors='replace')
    rdr = csv.DictReader(io.StringIO(text))
    fields = {f.strip().lower(): f for f in rdr.fieldnames or []}
    def f(*names):
        for n in names:
            if n in fields:
                return fields[n]
        raise ValueError(f'Usage CSV missing column: {names[0]}')
    c_name = next((fields[n] for n in ('name', 'hostname', 'host name') if n in fields), None)   # optional: blank when absent
    c_mac = f('mac address', 'mac')
    c_up, c_down, c_tot = f('upload (mb)', 'upload'), f('download (mb)', 'download'), f('total (mb)', 'total')
    c_pct = fields.get('% usage')
    usage = collections.OrderedDict()
    for row in rdr:
        mac = (row.get(c_mac) or '').strip().upper().replace('-', ':')
        if not mac:
            continue
        d = usage.setdefault(mac, {'name': ((row.get(c_name) if c_name else '') or '').strip(), 'up': 0.0, 'down': 0.0, 'tot': 0.0, 'pct': 0.0,
                                   'is_mac': bool(MAC_RE.match(mac))})
        d['up'] += float(row.get(c_up) or 0)
        d['down'] += float(row.get(c_down) or 0)
        d['tot'] += float(row.get(c_tot) or 0)
        if c_pct:
            d['pct'] += float((row.get(c_pct) or '0').replace('%', '') or 0)
    return usage


def merge(event_fileobj, usage_fileobj):
    logins, last = parse_event_log(event_fileobj)
    usage = parse_usage(usage_fileobj)
    rows = []
    for mac, d in usage.items():
        c = logins.get(mac)
        if c:
            primary = c.most_common(1)[0][0]
            users = ', '.join(f'{u} ({n})' for u, n in c.most_common())
            lt, lu = (last[mac][0].strftime('%Y-%m-%d %H:%M'), last[mac][1]) if mac in last else ('', '')
            n = sum(c.values())
        else:
            primary = users = lt = lu = ''
            n = 0
        rows.append({
            'user': primary, 'mac': mac, 'name': d['name'],
            'up': round(d['up'], 2), 'down': round(d['down'], 2),
            'tot': round(d['tot'], 2), 'gb': round(d['tot'] / 1024, 2),
            'logins': n, 'users': users, 'last_user': lu, 'last_time': lt,
            'pct': round(d['pct'], 2), 'is_mac': d['is_mac'],
        })
    rows.sort(key=lambda r: -r['tot'])
    per_user = collections.OrderedDict()
    for r in rows:
        if r['user']:
            p = per_user.setdefault(r['user'], {'macs': [], 'devs': [], 'up': 0.0, 'down': 0.0, 'tot': 0.0})
            p['macs'].append(r['mac'])
            p['devs'].append(dev_label(r['mac'], r['name']))
            p['up'] += r['up']
            p['down'] += r['down']
            p['tot'] += r['tot']
    per_user = collections.OrderedDict(sorted(per_user.items(), key=lambda x: -x[1]['tot']))
    unmatched = [r for r in rows if not r['user']]
    non_mac = sum(1 for r in rows if not r['is_mac'])
    summary = {
        'non_mac': non_mac,
        'devices': len(rows), 'matched': len(rows) - len(unmatched), 'unmatched': len(unmatched),
        'users': len(per_user), 'login_macs': len(logins),
        'total_gb': round(sum(r['tot'] for r in rows) / 1024, 2),
        'up_gb': round(sum(r['up'] for r in rows) / 1024, 2), 'down_gb': round(sum(r['down'] for r in rows) / 1024, 2),
        'unmatched_gb': round(sum(r['tot'] for r in unmatched) / 1024, 2),
    }
    return rows, per_user, unmatched, summary


def build_xlsx(rows, per_user, unmatched):
    out = openpyxl.Workbook()
    s = out.active
    s.title = 'Device usage by user'
    s.append(['User ID (primary)', 'MAC address', 'Hostname', 'Upload (MB)', 'Download (MB)',
              'Total (MB)', 'Total (GB)', '% Usage', 'Logins', 'All user IDs seen (logins)', 'Last login user', 'Last login time'])
    for r in rows:
        s.append([r['user'], r['mac'], r['name'], r['up'], r['down'], r['tot'], r['gb'], r['pct'],
                  r['logins'], r['users'], r['last_user'], r['last_time']])
    s2 = out.create_sheet('Usage per user')
    s2.append(['User ID', 'Devices (MAC (hostname))', 'No. of devices', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Total (GB)'])
    for u, p in per_user.items():
        s2.append([u, ', '.join(p['devs']), len(p['macs']), round(p['up'], 2), round(p['down'], 2), round(p['tot'], 2), round(p['tot'] / 1024, 2)])
    s3 = out.create_sheet('Unmatched devices')
    s3.append(['MAC address', 'Hostname', 'Upload (MB)', 'Download (MB)', 'Total (MB)', 'Note'])
    for r in unmatched:
        s3.append([r['mac'], r['name'], r['up'], r['down'], r['tot'], 'No login found' if r['is_mac'] else 'Not a MAC (IP / non-portal traffic)'])
    for sh in out:
        for c in sh[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1F4E78')
        sh.freeze_panes = 'A2'
        sh.auto_filter.ref = sh.dimensions
        for col in sh.columns:
            w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
            sh.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, w + 2), 60)
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = '#,##0.00'
    buf = io.BytesIO()
    out.save(buf)
    buf.seek(0)
    return buf

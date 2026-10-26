"""Core behaviour: attribution rules, import / delete rules, report views, Excel layout, security headers, audit trail, backup.
Run:  python -m unittest discover -s tests      (scratch databases only)"""
import io, os, unittest
from datetime import date, datetime

import openpyxl
from support import A, B, events_xlsx, flashes, load, signed_in, usage_csv


class Attribution(unittest.TestCase):
    def setUp(self):
        self.app, self.store = load()
        import daily
        self.daily = daily

    def test_same_day_login_wins_with_most_logins(self):
        ev = [(datetime(2026, 10, 3, 9), 'ana'), (datetime(2026, 10, 3, 10), 'bob'), (datetime(2026, 10, 3, 11), 'bob')]
        user, basis, _ = self.daily.attribute(ev, date(2026, 10, 3))
        self.assertEqual((user, basis), ('bob', 'login'))

    def test_carried_from_earlier_login(self):
        user, basis, _ = self.daily.attribute([(datetime(2026, 10, 1, 9), 'ana')], date(2026, 10, 5))
        self.assertEqual((user, basis), ('ana', 'carried'))

    def test_inferred_from_next_login_when_log_starts_later(self):
        user, basis, _ = self.daily.attribute([(datetime(2026, 10, 6, 9), 'bob')], date(2026, 10, 3))
        self.assertEqual((user, basis), ('bob', 'inferred'))

    def test_no_login_at_all(self):
        self.assertEqual(self.daily.attribute([], date(2026, 10, 3))[:2], ('', ''))

    def test_period_total_equals_sum_of_days_and_gb_is_mb_over_1024(self):
        ev = {A: [(datetime(2026, 10, 1, 8), 'ana')]}
        byday = [(date(2026, 10, 1), {A: dict(name='h', up=10, down=90, tot=100, is_mac=True)}),
                 (date(2026, 10, 2), {A: dict(name='h', up=20, down=180, tot=200, is_mac=True)})]
        detail, pivot, days, summary = self.daily.compute(ev, byday)
        self.assertEqual(pivot[0]['total'], 300)
        self.assertEqual((pivot[0]['up'], pivot[0]['down']), (30, 270))
        self.assertEqual(pivot[0]['gb'], round(300 / 1024, 2))
        self.assertEqual(summary['total_gb'], round(300 / 1024, 2))


class Imports(unittest.TestCase):
    def setUp(self):
        self.app, self.store = load()

    def test_reimport_same_day_replaces_not_adds(self):
        for _ in range(2):
            self.store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(usage_csv(('h', A, 1, 9))))
        ev, byday = self.store.load_for_report('S', date(2026, 10, 1), date(2026, 10, 1))
        self.assertEqual(byday[0][1][A]['tot'], 10)

    def test_empty_usage_file_is_refused_and_keeps_the_day(self):
        self.store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(usage_csv(('h', A, 1, 9))))
        with self.assertRaises(ValueError):
            self.store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(usage_csv()))
        self.assertEqual(self.store.status('S')['days'], ['2026-10-01'])
        with self.assertRaises(ValueError):
            self.store.import_period('S', 'p_2026-10-01_2026-10-07.csv', io.BytesIO(usage_csv()))

    def test_thousands_separators_are_accepted(self):
        csv = (b'Name,MAC Address,Upload (MB),Download (MB),Total (MB)\nh,' + A.encode() + b',"1,234.5","2,000.0","3,234.5"\n')
        self.store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(csv))
        _, byday = self.store.load_for_report('S', date(2026, 10, 1), date(2026, 10, 1))
        self.assertAlmostEqual(byday[0][1][A]['tot'], 3234.5)

    def test_real_non_number_gets_a_readable_error(self):
        csv = b'Name,MAC Address,Upload (MB),Download (MB),Total (MB)\nh,' + A.encode() + b',abc,1,1\n'
        with self.assertRaisesRegex(ValueError, 'not a number'):
            self.store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(csv))

    def test_missing_hostname_column_is_ok(self):
        csv = b'MAC Address,Upload (MB),Download (MB),Total (MB)\n' + A.encode() + b',1,2,3\n'
        self.store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(csv))
        _, byday = self.store.load_for_report('S', date(2026, 10, 1), date(2026, 10, 1))
        self.assertEqual(byday[0][1][A]['name'], '')

    def test_events_deduplicated_and_unreadable_dates_counted(self):
        f = lambda: events_xlsx((datetime(2026, 10, 1, 8), A, 'ana'), ('not a date', A, 'ana'))
        first = self.store.import_events('S', 'e.xlsx', f())
        self.assertEqual((first['rows'], first['added'], first['skipped']), (1, 1, 1))
        again = self.store.import_events('S', 'e.xlsx', f())
        self.assertEqual(again['added'], 0)

    def test_period_reversed_dates_are_swapped(self):
        r = self.store.import_period('S', 'p_2026-10-09_2026-10-03.csv', io.BytesIO(usage_csv(('h', A, 1, 2))))
        self.assertEqual((r['start'], r['end']), (date(2026, 10, 3), date(2026, 10, 9)))


class DeleteAndAudit(unittest.TestCase):
    def setUp(self):
        self.app, self.store = load()
        for d in ('01', '02', '03'):
            self.store.import_usage('S', f'x_2026-10-{d}.csv', io.BytesIO(usage_csv(('h', A, 1, 2))), user='admin')
        self.store.import_events('S', 'e.xlsx', events_xlsx((datetime(2026, 10, 1, 8), A, 'ana'), (datetime(2026, 10, 3, 8), A, 'ana')), user='admin')

    def test_usage_delete_respects_the_date_range(self):
        n = self.store.delete_data('S', 'usage', date(2026, 10, 2), date(2026, 10, 3), user='bob')
        self.assertEqual(n, 2)
        self.assertEqual(self.store.status('S')['days'], ['2026-10-01'])

    def test_logins_delete_respects_the_range(self):
        self.store.delete_data('S', 'logins', date(2026, 10, 1), date(2026, 10, 2))
        self.assertEqual(self.store.status('S')['logins'], 1)

    def test_everything_needs_the_site_name_typed(self):
        c, t = signed_in(self.app)
        c.post('/delete', data={'_csrf': t, 'site': 'S', 'what': 'all', 'confirm': 'nope'})
        self.assertEqual(len(self.store.status('S')['days']), 3)
        c.post('/delete', data={'_csrf': t, 'site': 'S', 'what': 'all', 'confirm': 's'})
        self.assertEqual(self.store.status('S')['days'], [])

    def test_audit_trail_records_who_imported_and_deleted(self):
        self.store.delete_data('S', 'usage', date(2026, 10, 1), date(2026, 10, 1), user='bob')
        rows = self.store.status('S')['imports']
        self.assertEqual(rows[0][1], 'delete')
        self.assertEqual(rows[0][7], 'bob')
        self.assertIn('admin', [r[7] for r in rows])

    def test_old_database_without_user_column_is_migrated_in_place(self):
        import sqlite3
        db = self.store.DB_PATH
        con = sqlite3.connect(db)
        con.executescript('DROP TABLE imports; CREATE TABLE imports (id INTEGER PRIMARY KEY AUTOINCREMENT, imported_at TEXT, site TEXT, kind TEXT, filename TEXT, rows INTEGER, added INTEGER, first_ts TEXT, last_ts TEXT);')
        con.commit(); con.close()
        self.store._ready = False
        self.assertEqual(len(self.store.status('S')['days']), 3)       # data kept
        self.store.import_usage('S', 'x_2026-10-09.csv', io.BytesIO(usage_csv(('h', A, 1, 2))), user='admin')
        self.assertEqual(self.store.status('S')['imports'][0][7], 'admin')


class ViewsAndExcel(unittest.TestCase):
    def setUp(self):
        self.app, self.store = load()
        self.store.import_events('S', 'e.xlsx', events_xlsx((datetime(2026, 10, 1, 8), A, 'ana')))
        for d in ('01', '02'):
            self.store.import_usage('S', f'x_2026-10-{d}.csv', io.BytesIO(usage_csv(('laptop', A, 10, 90), ('', B, 1, 2), ('ip', '10.0.0.9', 1, 1))))
        self.c, self.t = signed_in(self.app)
        self.q = '/report?site=S&start=2026-10-01&end=2026-10-02'

    def sheet(self, url, name):
        wb = openpyxl.load_workbook(io.BytesIO(self.c.get(url).data))
        return [[c.value for c in r] for r in wb[name]]

    def test_total_period_is_default_and_has_upload_download(self):
        html = self.c.get(self.q).get_data(as_text=True)
        self.assertIn('GB uploaded', html)
        self.assertIn('Upload GB', html)           # unmatched shows MB and GB

    def test_total_excel_layout_macs_last_and_unmatched_gb(self):
        users = self.sheet(self.q + '&view=total&download=1', 'Usage per user')
        self.assertEqual(users[0][-1], 'MACs (hostname)')
        self.assertEqual(users[1][3], 180)         # download MB over two days
        unm = self.sheet(self.q + '&view=total&download=1', 'Unmatched')
        self.assertIn('Total (GB)', unm[0])

    def test_day_excel_has_metric_sheets_and_device_download_follows_view(self):
        wb = openpyxl.load_workbook(io.BytesIO(self.c.get(self.q + '&view=day&download=1').data))
        for name in ('Usage per user per day', 'User per day - Upload', 'User per day - Download', 'Device per day - Download', 'Device-day detail'):
            self.assertIn(name, wb.sheetnames)
        dev_total = openpyxl.load_workbook(io.BytesIO(self.c.get(self.q + '&view=total&download=device').data))
        self.assertEqual(dev_total.sheetnames, ['Usage per device'])

    def test_period_page_and_downloads(self):
        self.store.import_period('S', 'p_2026-10-01_2026-10-07.csv', io.BytesIO(usage_csv(('laptop', A, 10, 90), ('', B, 1, 2))))
        url = '/period?site=S&start=2026-10-01&end=2026-10-07'
        self.assertIn('GB downloaded', self.c.get(url).get_data(as_text=True))
        wb = openpyxl.load_workbook(io.BytesIO(self.c.get(url + '&download=1').data))
        self.assertEqual(wb.sheetnames[0], 'Usage per device')
        self.assertEqual(self.c.get(url + '&download=device').status_code, 200)

    def test_bad_inputs_redirect_instead_of_crashing(self):
        for u in ('/report?site=S&start=2026-10-05&end=2026-10-01', '/report?site=S&start=garbage', '/period?site=S&start=x&end=y', '/report?site=NOPE'):
            self.assertIn(self.c.get(u).status_code, (200, 302), u)


class Security(unittest.TestCase):
    def test_signed_in_pages_are_not_cacheable_and_have_hardening_headers(self):
        app, _ = load()
        c, _t = signed_in(app)
        r = c.get('/')
        self.assertEqual(r.headers['Cache-Control'], 'no-store')
        self.assertEqual(r.headers['X-Content-Type-Options'], 'nosniff')
        self.assertEqual(r.headers['X-Frame-Options'], 'SAMEORIGIN')
        r = c.get('/static/kck-ui/kck-ui.css')
        self.assertNotEqual(r.headers.get('Cache-Control'), 'no-store')
        r.close()

    def test_login_required_and_health_open(self):
        app, _ = load()
        c = app.app.test_client()
        self.assertEqual(c.get('/').status_code, 302)
        self.assertEqual(c.get('/health').status_code, 200)
        self.assertEqual(c.post('/delete', data={'site': 'S', 'what': 'all'}).status_code, 302)

    def test_extra_accounts_and_csrf(self):
        app, _ = load(ADMIN_USERS='bob:pw2')
        c, t = signed_in(app, 'bob', 'pw2')
        self.assertEqual(c.get('/').status_code, 200)
        c.post('/delete', data={'site': 'S', 'what': 'usage'})        # no CSRF token: refused
        self.assertIn('expired', ' '.join(flashes(c)))

    def test_lockout_after_five_failures_and_per_client_with_proxy(self):
        app, _ = load()
        c = app.app.test_client()
        for _ in range(6):
            c.get('/login')
            with c.session_transaction() as s:
                t = s['csrf']
            r = c.post('/login', data={'_csrf': t, 'username': 'admin', 'password': 'bad'})
        self.assertEqual(r.status_code, 429)
        # behind one proxy each client is counted separately: client 2 is not locked out by client 1
        app2, _ = load(TRUSTED_PROXIES='1')
        c2 = app2.app.test_client()
        def attempt(ip):
            c2.get('/login', headers={'X-Forwarded-For': ip})
            with c2.session_transaction() as s:
                t = s['csrf']
            return c2.post('/login', data={'_csrf': t, 'username': 'admin', 'password': 'bad'}, headers={'X-Forwarded-For': ip}).status_code
        for _ in range(6):
            last = attempt('10.1.1.1')
        self.assertEqual(last, 429)
        self.assertEqual(attempt('10.2.2.2'), 200)

    def test_old_lockout_entries_are_pruned(self):
        app, _ = load()
        app._fails.update({f'1.1.1.{i}': [1, 0, 0] for i in range(50)})
        app._prune_fails(10_000_000)
        self.assertEqual(app._fails, {})


class Backup(unittest.TestCase):
    def test_backup_is_a_usable_copy_and_old_ones_are_pruned(self):
        import sqlite3, tempfile, time
        app, store = load()
        store.import_usage('S', 'x_2026-10-01.csv', io.BytesIO(usage_csv(('h', A, 1, 2))))
        d = tempfile.mkdtemp()
        for _ in range(3):
            out = app.backup_db(d, keep=2)
            time.sleep(1.05)
        self.assertEqual(len(os.listdir(d)), 2)
        con = sqlite3.connect(out)
        n = con.execute('SELECT COUNT(*) FROM usage_daily').fetchone()[0]
        con.close()
        self.assertEqual(n, 1)


if __name__ == '__main__':
    unittest.main()
